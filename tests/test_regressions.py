"""User-facing regressions; run with python -m unittest discover -s tests."""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills/ppt-master/scripts"
sys.path.insert(0, str(SCRIPTS))
TEMP = ROOT / ".cache" / "tests"
TEMP.mkdir(parents=True, exist_ok=True)

from template_fill_pptx.scaffolder import scaffold_plan
from review_pptx import Box, check_bounds, review
from pptx import Presentation
from pptx.util import Inches


class ScaffoldTests(unittest.TestCase):
    def setUp(self):
        self.library = {"slides": [{"slide_index": i, "slots": [
            {"slot_id": "title", "text": "Example"}]} for i in (1, 2, 3)]}

    def test_order_and_repeated_pages(self):
        plan = scaffold_plan(self.library, [3, 1, 3])
        self.assertEqual([s["source_slide"] for s in plan["slides"]], [3, 1, 3])
        plan["slides"][0]["replacements"][0]["text"] = "Changed"
        self.assertEqual(plan["slides"][2]["replacements"][0]["text"], "Example")

    def test_unknown_page_fails_instead_of_silently_dropping(self):
        with self.assertRaisesRegex(RuntimeError, "99"):
            scaffold_plan(self.library, [1, 99])

    def test_default_pages(self):
        self.assertEqual(len(scaffold_plan(self.library)["slides"]), 3)

    def test_repeated_chart_edits_are_independent(self):
        library = {"slides": [{"slide_index": 1, "charts": [{
            "chart_id": "chart1", "categories": ["A"],
            "series": [{"name": "S", "values": [1]}]}]}]}
        plan = scaffold_plan(library, [1, 1])
        first = plan["slides"][0]["chart_edits"][0]
        first["categories"][0] = "Changed"
        first["series"][0]["values"][0] = 9
        second = plan["slides"][1]["chart_edits"][0]
        self.assertEqual(second["categories"], ["A"])
        self.assertEqual(second["series"][0]["values"], [1])
        self.assertEqual(library["slides"][0]["charts"][0]["categories"], ["A"])


class ExportTests(unittest.TestCase):
    def test_failed_export_has_nonzero_exit(self):
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            result = subprocess.run([sys.executable, str(SCRIPTS / "svg_to_pptx.py"),
                                     str(Path(tmp) / "missing")], capture_output=True)
        self.assertNotEqual(result.returncode, 0)


class ReviewTests(unittest.TestCase):
    def test_vertical_bounds(self):
        box = Box(1, 1, "title", "text", 0, 0, 100, 100, True,
                  text_emu=(0, 1000000, 100000, 100000))
        self.assertIsNotNone(check_bounds(box, 2000000, 500000))

    def test_placeholders_in_groups_and_tables(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = "{{姓名}}"
        table = slide.shapes.add_table(1, 1, 0, Inches(2), Inches(3), Inches(1)).table
        table.cell(0, 0).text = "{{result}}"
        with tempfile.TemporaryDirectory(dir=TEMP) as tmp:
            path = Path(tmp) / "review.pptx"
            prs.save(path)
            issues, _ = review(str(path))
        self.assertEqual(sum(i.kind == "UNREPLACED" for i in issues), 2)


if __name__ == "__main__":
    unittest.main()
