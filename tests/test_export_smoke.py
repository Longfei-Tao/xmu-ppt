"""Export every shipped style, preserving native text and embedded pictures."""
import contextlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/ppt-master/scripts"))
from svg_to_pptx import create_pptx_with_native_svg
from template_fill_pptx.applier import apply_plan
from template_fill_pptx.scaffolder import scaffold_plan
from pptx import Presentation

DECKS = ROOT / "skills/ppt-master/templates/decks"


class ExportSmokeTests(unittest.TestCase):
    def test_all_svg_styles_export_editable_text_and_images(self):
        cache = ROOT / ".cache"
        cache.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=cache) as tmp:
            for deck in sorted(DECKS.glob("厦门大学_*")):
                with self.subTest(deck=deck.name), contextlib.redirect_stdout(io.StringIO()):
                    target = Path(tmp) / (deck.name + ".pptx")
                    self.assertTrue(create_pptx_with_native_svg(
                        sorted(deck.glob("*.svg")), target, canvas_format="ppt169",
                        verbose=False, use_native_shapes=True, workers=1))
                    prs = Presentation(target)
                    self.assertEqual(len(prs.slides), 5)
                    with zipfile.ZipFile(target) as archive:
                        self.assertIsNone(archive.testzip())
                        slides = [archive.read(f"ppt/slides/slide{i}.xml") for i in range(1, 6)]
                        self.assertTrue(all(b"<a:t>" in s for s in slides))
                        self.assertTrue(any(n.startswith("ppt/media/") for n in archive.namelist()))

    def test_native_styles_repeat_reorder_and_fill(self):
        cache = ROOT / ".cache"
        cache.mkdir(exist_ok=True)
        with tempfile.TemporaryDirectory(dir=cache) as tmp:
            for source in sorted(DECKS.glob("*/source.pptx")):
                with self.subTest(deck=source.parent.name):
                    library = json.loads(source.with_suffix(".slide_library.json").read_text(encoding="utf-8"))
                    plan = scaffold_plan(library, [2, 1, 2])
                    for index, page in enumerate(plan["slides"]):
                        for slot in page["replacements"]:
                            slot["text"] = f"验收{index}"
                    target = Path(tmp) / (source.parent.name + ".pptx")
                    apply_plan(source, plan, target)
                    prs = Presentation(target)
                    self.assertEqual(len(prs.slides), 3)
                    for index, slide in enumerate(prs.slides):
                        self.assertIn(f"验收{index}", slide._element.xml)
                    with zipfile.ZipFile(target) as archive:
                        self.assertIsNone(archive.testzip())
                        self.assertTrue(any(n.startswith("ppt/media/") for n in archive.namelist()))
